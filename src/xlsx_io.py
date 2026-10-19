"""Read every model input and parameter back from the workbook so the Python engine runs on exactly what the Excel shows.
The workbook must have cached values (saved by Excel, or recalculated with LibreOffice)."""
from openpyxl import load_workbook
from layout import FIELDS
import engine as E

RATINGS = ['B+', 'B', 'B-', 'CCC+', 'CCC']

def _v(x):
    return 0.0 if x is None or x == '' else x

def read_inputs(path):
    wb = load_workbook(path, data_only=True)
    ib = wb['Inputs_Banks']; keys = [f[0] for f in FIELDS]
    ids, banks = [], {}
    r = 6
    while ib.cell(r, 1).value:
        row = {k: ib.cell(r, c + 1).value for c, k in enumerate(keys)}
        bid = row['id']; ids.append(bid)
        row['dg'] = [float(_v(row[f'dg{n}'])) for n in range(7)]
        for k in list(row):
            if k not in ('id', 'name', 'ccy', 'dg') and not (k.startswith('dg') and k[2:].isdigit()):
                row[k] = float(_v(row[k]))
        banks[bid] = row; r += 1
    h = wb['Holdings']
    classes = [h.cell(5, 3 + k).value for k in range(13)]
    holdings = {bid: {classes[k]: float(_v(h.cell(6 + i, 3 + k).value)) for k in range(13)} for i, bid in enumerate(ids)}
    CL = [[classes[k], float(h.cell(31, 3 + k).value), float(h.cell(32, 3 + k).value), float(h.cell(33, 3 + k).value),
           float(h.cell(34, 3 + k).value), float(_v(h.cell(39, 3 + k).value)), ''] for k in range(13)]
    MB = {classes[k]: float(_v(h.cell(36, 3 + k).value)) for k in range(13)}
    INS = {classes[k]: float(_v(h.cell(37, 3 + k).value)) for k in range(13)}
    b = wb['BDC']; bdc = []
    r = 6
    while b.cell(r, 1).value:
        bdc.append(dict(name=b.cell(r, 1).value, type=b.cell(r, 2).value, assets=float(b.cell(r, 3).value), sw=float(b.cell(r, 4).value),
                        de=float(b.cell(r, 5).value), pnav=b.cell(r, 6).value, unsec=float(b.cell(r, 7).value), redq=float(b.cell(r, 8).value),
                        hy=float(_v(b.cell(r, 20).value))))
        r += 1
    e = wb['BDC_EAD']
    veh = [e.cell(6 + v, 1).value for v in range(22) if e.cell(6 + v, 1).value]
    blk = {'spv_c': 6, 'spv_o': 31, 'rev_c': 56, 'rev_o': 81}
    ead = {k: {bid: [float(_v(e.cell(r0 + v, 3 + i).value)) for v in range(len(veh))] for i, bid in enumerate(ids)} for k, r0 in blk.items()}
    a = wb['AI_Deals']; deals = []
    r = 6
    hdrs = [a.cell(5, c).value for c in range(1, a.max_column + 1)]
    bcol = {bid: hdrs.index(bid) + 1 for bid in ids}
    while a.cell(r, 2).value and a.cell(r, 3).value:
        x = dict(idx=a.cell(r, 1).value, deal=a.cell(r, 2).value, cat=a.cell(r, 3).value, lessee=a.cell(r, 4).value or '',
                 ltv=float(_v(a.cell(r, 7).value)), frep=float(_v(a.cell(r, 8).value)), cash=float(_v(a.cell(r, 9).value)),
                 lgd=float(_v(a.cell(r, 10).value)), uw=float(_v(a.cell(r, 11).value)), senior=float(_v(a.cell(r, 12).value)),
                 mtm0=float(_v(a.cell(r, 13).value)), size=a.cell(r, 19).value,
                 holds={bid: float(_v(a.cell(r, bcol[bid]).value)) for bid in ids})
        deals.append(x); r += 1
    n = wb['Network']
    X = [[float(n.cell(6 + j, 3 + i).value) for j in range(len(ids))] for i in range(len(ids))]   # X[i][j], sheet is transposed
    g = wb['Globals']; G = {}
    r = 5
    while g.cell(r, 1).value:
        G[g.cell(r, 1).value] = float(g.cell(r, 2).value); r += 1
    s = wb['Scenarios']; SC = []
    while s.cell(4, 5 + len(SC)).value not in (None, 'ACTIVE'): SC.append(s.cell(4, 5 + len(SC)).value)
    Pm = {}
    r = 5
    while s.cell(r, 1).value:
        Pm[s.cell(r, 1).value] = [float(_v(s.cell(r, 5 + j).value)) for j in range(len(SC))]; r += 1
    ce = wb['CreditEngine']
    POOL_W = [float(ce.cell(6 + k, 2).value) for k in range(5)]; PD3 = [float(ce.cell(6 + k, 3).value) for k in range(5)]
    TR_AVG = {RATINGS[k]: [float(ce.cell(6 + k, 8).value), float(ce.cell(6 + k, 9).value)] for k in range(5)}
    TR_09 = {RATINGS[k]: [float(ce.cell(6 + k, 10).value), float(ce.cell(6 + k, 11).value)] for k in range(5)}
    LESSEE = {}
    for r in range(1, ce.max_row + 1):
        if ce.cell(r, 1).value == 'Key' and ce.cell(r, 2).value == 'PD 3-yr':
            rr = r + 1
            while ce.cell(rr, 1).value:
                key = ce.cell(rr, 1).value
                LESSEE[key] = None if key == 'ANTH' else float(ce.cell(rr, 2).value); rr += 1
            break
    cfg = dict(G=G, CL=CL, MB=MB, INS=INS, POOL_W=POOL_W, PD3=PD3, TR_AVG=TR_AVG, TR_09=TR_09, LESSEE=LESSEE, SCEN=SC, P=Pm)
    d = dict(ids=ids, banks=banks, holdings=holdings, bdc=bdc, bdc_veh=veh, bdc_ead=ead, deals=deals, network=dict(X=X), cfg=cfg)
    return E.prep(d)
