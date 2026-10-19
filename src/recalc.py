"""Recalculate every formula in an .xlsx with LibreOffice and report formula errors as JSON.

Usage:  python src/recalc.py workbook.xlsx [timeout_seconds]

openpyxl writes formulas without cached values, so a workbook built by build_xlsx.py shows empty cells
in viewers (and in xlsx_io.read_inputs) until something recalculates it. Excel does that on open;
this script does it headless so the Python side can read the results back.
Prints {"status": "success" | "errors_found", "total_formulas": n, "total_errors": n, "errors": {...}}.
Set SOFFICE to the LibreOffice binary if it is not on PATH (e.g. C:/Program Files/LibreOffice/program/soffice.exe).
"""
import json, os, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub Recalc()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""
ERRORS = ('#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A')


def recalc(path, timeout=300):
    soffice = os.environ.get('SOFFICE', 'soffice')
    env = dict(os.environ, SAL_USE_VCLPLUGIN='svp')
    with tempfile.TemporaryDirectory(prefix='lo_profile_') as prof:
        uri = Path(prof).as_uri()
        # first start creates the profile; then drop the macro into the Standard library
        subprocess.run([soffice, '--headless', '--terminate_after_init', f'-env:UserInstallation={uri}'], env=env, capture_output=True, timeout=120)
        lib = Path(prof) / 'user' / 'basic' / 'Standard'
        lib.mkdir(parents=True, exist_ok=True)
        (lib / 'Module1.xba').write_text(MACRO)
        subprocess.run([soffice, '--headless', '--norestore', f'-env:UserInstallation={uri}',
                        'vnd.sun.star.script:Standard.Module1.Recalc?language=Basic&location=application', str(Path(path).resolve())],
                       env=env, capture_output=True, timeout=timeout)
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    n, errs = 0, {}
    for ws in formulas.worksheets:
        wv = values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    n += 1
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v in ERRORS:
                        errs.setdefault(v, []).append(f'{ws.title}!{c.coordinate}')
    total = sum(len(x) for x in errs.values())
    return dict(status='errors_found' if total else 'success', total_formulas=n, total_errors=total,
                errors={k: dict(count=len(x), locations=x[:50]) for k, x in errs.items()})


if __name__ == '__main__':
    print(json.dumps(recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 300), indent=2))
