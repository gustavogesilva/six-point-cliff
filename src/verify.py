"""Check that the Excel formulas reproduce the Python engine in every scenario.

Usage (from the repo root):  python src/verify.py [workbook.xlsx]
Sets each scenario on Dashboard!C4 in a temporary copy, recalculates it with LibreOffice and prints the Checks sheet.
"""
import json, os, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook
from params import SCEN
ROOT = Path(__file__).resolve().parents[1]
RECALC = os.environ.get('RECALC_SCRIPT', str(ROOT / 'src' / 'recalc.py'))
SRC = sys.argv[1] if len(sys.argv) > 1 else str(ROOT / 'model' / 'Private_Credit_Contagion_Model.xlsx')
T = os.path.join(tempfile.mkdtemp(), 'verify.xlsx')
for s in SCEN:
    wb = load_workbook(SRC); wb['Dashboard']['C4'] = s; wb.save(T)
    r = json.loads(subprocess.run(['python3', RECALC, T, '240'], capture_output=True, text=True).stdout)
    v = load_workbook(T, data_only=True); ck = v['Checks']; res = v['Results']
    checks = [(ck.cell(rr, 1).value[:40], ck.cell(rr, 2).value, ck.cell(rr, 4).value) for rr in range(5, 12)]
    sysv = [res.cell(r_, 2).value for r_ in range(33, 46)]
    print(f"{s:22s} recalc={r['status']} errs={r['total_errors']} | ratio diff={checks[5][1]:.2e} {checks[5][2]} | loss diff={checks[6][1]:.2e} {checks[6][2]} | other checks: {[c[2] for c in checks[:5]]}")
    print('    system:', sysv[:10])
